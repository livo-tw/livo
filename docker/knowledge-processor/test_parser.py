import base64
import importlib.util
import io
import json
import os
import re
import socket
import subprocess
import sys
import threading
import time
import unittest
import urllib.error
import urllib.request
import zipfile
from http.server import ThreadingHTTPServer
from pathlib import Path
from unittest.mock import patch

sys.path.insert(0, str(Path(__file__).parent))
from parser import ImportFailure, checked_zip, markdown, parse, sanitize
import server


def docx_fixture():
    from docx import Document
    doc = Document()
    doc.add_heading('Project Alpha', 1)
    doc.add_paragraph('A reviewed decision with <script>literal text</script>.')
    table = doc.add_table(rows=1, cols=2)
    table.cell(0, 0).text = 'Item'
    table.cell(0, 1).text = 'Decision'
    out = io.BytesIO()
    doc.save(out)
    return out.getvalue()


def pdf_fixture(scanned=False, mixed=False):
    from pypdf import PdfWriter
    from pypdf.generic import DictionaryObject, NameObject, NumberObject, DecodedStreamObject
    writer = PdfWriter()
    page = writer.add_blank_page(600, 180)
    font = DictionaryObject({NameObject('/Type'): NameObject('/Font'), NameObject('/Subtype'): NameObject('/Type1'), NameObject('/BaseFont'): NameObject('/Helvetica')})
    page[NameObject('/Resources')] = DictionaryObject({NameObject('/Font'): DictionaryObject({NameObject('/F1'): writer._add_object(font)})})
    stream = DecodedStreamObject()
    stream.set_data(b'BT /F1 20 Tf 15 130 Td (Project Alpha has a confirmed text decision ready for review.) Tj ET')
    page[NameObject('/Contents')] = writer._add_object(stream)
    if scanned or mixed:
        if scanned: writer.remove_page(0)
        glyphs = {'A':['01110','10001','10001','11111','10001','10001','10001'], 'L':['10000']*6+['11111'], 'P':['11110','10001','10001','11110','10000','10000','10000'], 'H':['10001']*3+['11111']+['10001']*3}
        text = 'ALPHA ALPHA ALPHA'
        scale = 10
        width, height = len(text)*6*scale+40, 130
        pixels = bytearray([255])*(width*height*3)
        for at, char in enumerate(text):
            for y, row in enumerate(glyphs.get(char, ['00000']*7)):
                for x, ink in enumerate(row):
                    if ink == '1':
                        for dy in range(scale):
                            for dx in range(scale):
                                pos = ((20+y*scale+dy)*width+20+at*6*scale+x*scale+dx)*3
                                pixels[pos:pos+3] = b'\x00\x00\x00'
        image = DecodedStreamObject()
        image.set_data(bytes(pixels))
        image.update({NameObject('/Type'):NameObject('/XObject'),NameObject('/Subtype'):NameObject('/Image'),NameObject('/Width'):NumberObject(width),NameObject('/Height'):NumberObject(height),NameObject('/ColorSpace'):NameObject('/DeviceRGB'),NameObject('/BitsPerComponent'):NumberObject(8)})
        page = writer.add_blank_page(width,height)
        page[NameObject('/Resources')] = DictionaryObject({NameObject('/XObject'):DictionaryObject({NameObject('/Scan'):writer._add_object(image.flate_encode())})})
        stream = DecodedStreamObject()
        stream.set_data(f'q {width} 0 0 {height} 0 0 cm /Scan Do Q'.encode())
        page[NameObject('/Contents')] = writer._add_object(stream)
    out=io.BytesIO()
    writer.write(out)
    return out.getvalue()


def payload(source, data):
    return {'source':source,'data':base64.b64encode(data).decode()}


class ParserTests(unittest.TestCase):
    def test_markdown_is_passive(self):
        body,warnings=markdown(b'# Decision\n<script>alert(1)</script>\n![private](https://example.com/pixel)\n- [x] old\n[bad](javascript:evil)')
        self.assertNotIn('<script>',body)
        self.assertNotIn('<img',body)
        self.assertIn('historical_checkboxes',warnings)
        self.assertNotIn('href="javascript:',body)
    def test_sanitizer_removes_active_attributes(self):
        clean=sanitize('<p onclick="x">Okay<script>secret</script><a href="javascript:x">bad</a></p>')
        self.assertEqual(clean,'<p>Okay<a>bad</a></p>')
    def test_docx_order_and_table(self):
        result=parse(payload('docx',docx_fixture()))
        self.assertIn('<h1>Project Alpha</h1>',result['body'])
        self.assertIn('<table>',result['body'])
        self.assertIn('&lt;script&gt;',result['body'])
    def test_zip_traversal_and_expansion(self):
        for name,data in [('../evil',b'bad'),('word/document.xml',b'a'*1_000_000)]:
            out=io.BytesIO()
            with zipfile.ZipFile(out,'w',zipfile.ZIP_DEFLATED) as z:z.writestr(name,data)
            with self.assertRaises(ImportFailure):checked_zip(out.getvalue())
    def test_xml_entities_and_macros_rejected(self):
        for name,data in [('word/document.xml',b'<!DOCTYPE x>'),('word/vbaProject.bin',b'bad')]:
            out=io.BytesIO()
            with zipfile.ZipFile(out,'w') as z:
                z.writestr(name,data)
                if name!='word/document.xml':z.writestr('word/document.xml',b'<x/>')
            with self.assertRaises(ImportFailure):checked_zip(out.getvalue())
    def test_pdf_text_and_scan_pending(self):
        with patch.dict(os.environ,{'KNOWLEDGE_OCR_ENABLED':'0'}):
            result=parse(payload('pdf',pdf_fixture(mixed=True)))
        self.assertEqual([p['state'] for p in result['pages']],['text','ocr_pending'])
        self.assertTrue(result['incomplete'])
    def test_pdf_with_broken_xref_offset_is_still_parsed(self):
        broken=re.sub(rb'startxref\s+\d+',b'startxref\n9',pdf_fixture())
        with patch.dict(os.environ,{'KNOWLEDGE_OCR_ENABLED':'0'}):
            result=parse(payload('pdf',broken))
        self.assertEqual([p['state'] for p in result['pages']],['text'])
        self.assertIn('Project Alpha',result['body'])
    def test_pdf_encrypted_rejected(self):
        from pypdf import PdfWriter
        writer=PdfWriter();writer.add_blank_page(200,200);writer.encrypt('test-only');out=io.BytesIO();writer.write(out)
        with self.assertRaisesRegex(ImportFailure,'encrypted_pdf'):parse(payload('pdf',out.getvalue()))
    def test_invalid_and_empty_rejected(self):
        for value in [{'source':'docm','data':'YQ=='},{'source':'md','data':'%%%'}]:
            with self.assertRaises(ImportFailure):parse(value)


TOKEN='t'*40
REAL_WORK=server.work
# The caller must be another process: a forked parse child would otherwise keep
# a copy of the caller's socket open, so closing it would not disconnect.
ABANDONING_CLIENT='''import socket,sys,time
raw=sys.argv[3].encode()
client=socket.create_connection((sys.argv[1],int(sys.argv[2])))
client.sendall(b"POST /parse HTTP/1.1\\r\\nHost: x\\r\\nContent-Type: application/json\\r\\nAuthorization: Bearer "+sys.argv[4].encode()+b"\\r\\nContent-Length: "+str(len(raw)).encode()+b"\\r\\n\\r\\n"+raw)
time.sleep(float(sys.argv[5]))
'''


def slow_or_real_work(payload,pipe):
    if payload.get('slow'):
        time.sleep(60)
    REAL_WORK(payload,pipe)


class LocalServerTests(unittest.TestCase):
    """The real handler on an ephemeral port; parse runs in its child process."""
    def setUp(self):
        self.env=patch.dict(os.environ,{'KNOWLEDGE_PROCESSOR_TOKEN':TOKEN,'KNOWLEDGE_OCR_ENABLED':'0'})
        self.env.start()
        self.httpd=ThreadingHTTPServer(('127.0.0.1',0),server.Handler)
        self.httpd.daemon_threads=True
        threading.Thread(target=self.httpd.serve_forever,daemon=True).start()
        self.url='http://127.0.0.1:%d/parse'%self.httpd.server_address[1]
    def tearDown(self):
        self.httpd.shutdown();self.httpd.server_close();self.env.stop()
    def post(self,body,token=TOKEN):
        request=urllib.request.Request(self.url,json.dumps(body).encode(),headers={'Content-Type':'application/json','Authorization':'Bearer '+token})
        try:
            with urllib.request.urlopen(request,timeout=30) as response:return response.status,json.load(response)
        except urllib.error.HTTPError as error:return error.code,json.load(error)
    def test_missing_token_refuses_every_request(self):
        with patch.dict(os.environ,{'KNOWLEDGE_PROCESSOR_TOKEN':''}):
            self.assertEqual(self.post(payload('md',b'# Title'),''),(503,{'error':'processor_not_configured'}))
        self.assertEqual(self.post(payload('md',b'# Title'),'x'*40)[0],401)
        status,body=self.post(payload('md',b'# Title'))
        self.assertEqual(status,200);self.assertIn('<h1>Title</h1>',body['result']['body'])
    def test_abandoned_request_frees_the_only_slot(self):
        with patch.object(server,'work',slow_or_real_work):
            raw=json.dumps({**payload('md',b'# Slow'),'slow':True})
            host,port=self.httpd.server_address
            client=subprocess.Popen([sys.executable,'-c',ABANDONING_CLIENT,host,str(port),raw,TOKEN,'2'])
            time.sleep(1)
            self.assertEqual(self.post(payload('md',b'# Busy'))[0],429)
            self.assertEqual(client.wait(timeout=10),0)  # the caller gives up and disconnects
            deadline=time.monotonic()+10
            while True:
                status,body=self.post(payload('md',b'# Next'))
                if status!=429 or time.monotonic()>deadline:break
                time.sleep(0.2)
            self.assertEqual(status,200)
            self.assertIn('<h1>Next</h1>',body['result']['body'])
    def test_budget_bounds_a_parse(self):
        self.assertLess(server.BUDGET,95)  # knowledgeImport.ts waits 95 s for the processor
        import parser as module
        self.assertLessEqual(module.OCR_START_BUDGET+15+25,server.BUDGET)


class StartupTests(unittest.TestCase):
    def test_missing_token_does_not_exit(self):
        with socket.socket() as probe:
            probe.bind(('127.0.0.1',0));port=probe.getsockname()[1]
        env={**os.environ,'KNOWLEDGE_PROCESSOR_TOKEN':'','KNOWLEDGE_PROCESSOR_PORT':str(port),'PYTHONDONTWRITEBYTECODE':'1'}
        process=subprocess.Popen([sys.executable,str(Path(__file__).parent/'server.py')],env=env,stdout=subprocess.DEVNULL,stderr=subprocess.PIPE)
        try:
            time.sleep(1.5)
            self.assertIsNone(process.poll())  # a crash would make Docker restart it in a loop
            request=urllib.request.Request('http://127.0.0.1:%d/parse'%port,b'{}',headers={'Authorization':'Bearer '})
            with self.assertRaises(urllib.error.HTTPError) as raised:urllib.request.urlopen(request,timeout=10)
            self.assertEqual(raised.exception.code,503)
        finally:
            process.terminate();process.wait(timeout=10);process.stderr.close()


class HTTPTests(unittest.TestCase):
    @classmethod
    def setUpClass(cls):
        cls.url=os.environ.get('PROCESSOR_TEST_URL')
        if not cls.url:raise unittest.SkipTest('HTTP integration requires the private processor container')
    def send(self,source,data,token=None):
        request=urllib.request.Request(self.url+'/parse',json.dumps(payload(source,data)).encode(),headers={'Content-Type':'application/json','Authorization':'Bearer '+(token or os.environ['PROCESSOR_TEST_TOKEN'])})
        with urllib.request.urlopen(request,timeout=120) as response:return json.load(response)['result']
    def test_unauthorized_is_rejected(self):
        with self.assertRaises(urllib.error.HTTPError) as raised:self.send('md',b'# Test','incorrect')
        self.assertEqual(raised.exception.code,401)
    def test_all_formats_and_real_private_ocr(self):
        for source,data in [('md',b'# Text\n- [ ] Original'),('docx',docx_fixture()),('pdf',pdf_fixture())]:
            self.assertTrue(self.send(source,data)['body'])
        result=self.send('pdf',pdf_fixture(mixed=True))
        self.assertEqual([p['state'] for p in result['pages']],['text','needs_review'])
        self.assertIn('ALPHA',result['pages'][1]['text'].upper())
        self.assertTrue(result['needs_review'])
        self.assertFalse(result['incomplete'])


if __name__=='__main__':unittest.main()
